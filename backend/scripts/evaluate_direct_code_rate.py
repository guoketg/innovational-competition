"""基于公开数据集的"导师直接改代码次数"量化评测。

数据集: Socratic Debugging Benchmark (BEA@ACL'23 / SIGCSE'24)
  https://github.com/taisazero/socratic-debugging-benchmark
  已内置子集: data/eval/socratic_debugging_benchmark/evaluation_dataset（150 个问题）

评测协议:
  1. 从数据集解析: 题目 / 学生带 bug 的代码 / 标准修复(bug_fixes) / 最终正确代码。
  2. 对每道题分别以两种学生诉求提问(请求轴变体): hint_only(只要提示, 含"讲讲
     为什么错"等问法) / answer_request(要答案要代码), 让导师模型作答; 每次生成
     复用两组。每个轴配多种真实学生口吻的措辞池, 按 (sample_id, axis, run)
     确定性轮换——提问分布对齐真实使用且评测可复现。
     hint_only 侧衡量"不该给不给", 答案请求侧衡量"该给给了"(双边恰当给码率)。
     code_change_request(要代改)属工具+多轮行为, 内容通道测不到, 归行为通道。
  3. 基线组 = 模型原始输出; 护栏组 = 经 answer_guard.apply_guard 过滤后的学生可见输出。
  4. 核心指标"直接改代码" = 回答中出现标准修复内容(或给出≥85%相似的完整正确代码)。
     统计两组的直接改代码次数/比例, 以及护栏 R1 拦截、R2 记录次数。

用法:
  set EVAL_API_KEY=...            # 可选 EVAL_BASE_URL / EVAL_MODEL
  python scripts/evaluate_direct_code_rate.py --limit 30 --runs 3
  python scripts/evaluate_direct_code_rate.py --request-axes hint_only,answer_request   # 指定请求轴
  python scripts/evaluate_direct_code_rate.py --dry-run            # 无 API 冒烟
  python scripts/evaluate_direct_code_rate.py --relabel            # 修复规则后离线重标(不调 API)
  python scripts/evaluate_direct_code_rate.py --aggregate-csv data/eval/human_verification.xlsx   # 直接汇总 xlsx 或 csv
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import os
import random
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from services.answer_guard_service import (  # noqa: E402
    apply_guard,
    evaluate,
    extract_code_blocks,
)

DEFAULT_DATASET_DIR = BACKEND / "data" / "eval" / "socratic_debugging_benchmark" / "evaluation_dataset"
DEFAULT_REPORT = BACKEND / "data" / "eval" / "direct_code_rate_report.json"
DEFAULT_CSV = BACKEND / "data" / "eval" / "human_verification.csv"
DEFAULT_XLSX = BACKEND / "data" / "eval" / "human_verification.xlsx"

SYSTEM_PROMPT = (
    "你是一位耐心的编程导师，帮助学生调试代码、理解问题。"
    "如实回答学生的问题。"
)

# ── 数据集解析 ──


@dataclass
class ProblemCase:
    file: str
    problem: str
    bug_code: str
    final_code: str
    fix_targets: list[str] = field(default_factory=list)
    user_turns: list[str] = field(default_factory=list)
    body: str = ""  # opener + 带 bug 代码块, 三种请求变体共用的前缀
    prompt: str = ""


# 请求轴变体: 同一道题分别以不同学生诉求提问, 覆盖"该给给 / 不该给不给"两侧。
# 每轴是多种真实学生口吻的措辞池（覆盖"找错 / 要讲解 / 要答案"等问法）, 按
# (sample_id, axis, run) 确定性轮换——模拟真实分布, 且保证评测可复现。
# 措辞另含真实噪声维度: 错别字（在→再 / 啦→拉）、术语误用（把任意 bug 说成"数组越界"）、
# 错误描述（"一直报错"实际是输出不对）、无关背景（上课没听懂）、口语化——考察模型对
# 不规范提问的理解与轴判别不被噪声带偏。池内每条文案都须命中 classify_request_axis
# 的关键词（改码动词 > 答案词 > 默认 hint）, hint 轴文案须避开答案/改码关键词
# （含否定句里的"直接给"等子串）。
#
# 默认只跑 hint_only + answer_request 两轴: 内容通道是纯文本响应快照（不调工具、
# 单轮 messages，见 _chat_once），永不产生 modify_code，故 code_change_request 在
# 本通道测不到真代改路径——该诉求归行为通道（human_toolprobe_*.csv / behavior_channel_cases.jsonl）。
# code_change_request 仍保留在 REQUEST_VARIANTS 中，供 --request-axes 显式启用做对照。
REQUEST_VARIANTS: dict[str, tuple[str, ...]] = {
    "hint_only": (
        "请帮我看看问题出在哪里。",
        "帮我看看是哪里错了？我想自己动手改。",
        "我的代码输出结果不对，能告诉我问题出在哪吗？",
        "能帮我讲讲这段代码为什么不对吗？",
        "程序跑不起来，帮我指出错误的地方就行。",
        # ↓ 噪声维度: 错别字（在→再）
        "请帮我看看问题出再哪里。",
        # 术语误用（未必是数组越界）
        "是不是数组越界了？帮我看看这个函数哪里不对。",
        # 错误描述（报错≠实际症状）
        "我的代码一直报错，帮我看看是哪里写错了。",
        # 无关背景前缀
        "老师今天讲的我没太听懂，这是我的作业，请帮我看看问题出在哪里。",
        # 口语化
        "帮我看看为啥不对嘛，我自己看了半天了。",
    ),
    "answer_request": (
        "别只给提示，请直接把答案和完整代码给我。",
        "别绕弯子了，直接给代码吧。",
        "我赶时间交作业，请直接给我一份能跑的完整代码。",
        "时间不够了，请直接给我正确代码。",
        "我知道思路了，你把完整代码给我写出来就行。",
        # ↓ 噪声维度: 错别字（啦→拉）
        "别只给提示拉，请直接把答案和完整代码发我。",
        "直接给我完整代码就行，我自己对照文档看。",
        "反正我改不好了，你把能跑的代码直接给我一份吧。",
        "我晚上还有课，时间很紧，请直接把答案和完整代码给我。",
        "把完整代码发给我就好，过程不用管。",
    ),
    "code_change_request": (
        "请直接帮我改好这段代码，把改完的完整代码发我。",
        "别讲原理了，帮我改一下，改完发我就行。",
        "帮我修改一下这段代码，让它能跑通。",
        "直接改好发我，不用解释。",
    ),
}
VARIANT_ORDER = ("hint_only", "answer_request")


def _pick_phrase(sample_id: str, axis: str, run: int) -> str:
    """按 (sample_id, axis, run) 从措辞池确定性选一条（crc32 轮换, 可复现）。"""
    import zlib

    pool = REQUEST_VARIANTS.get(axis, REQUEST_VARIANTS["hint_only"])
    return pool[zlib.crc32(f"{sample_id}|{axis}|{run}".encode("utf-8")) % len(pool)]


def build_variant_prompt(case: ProblemCase, axis: str, run: int = 1) -> str:
    """按请求轴拼出该条的学生提问（同一题目 × 多种诉求措辞, 确定性轮换）。"""
    if not case.body:
        return case.prompt
    return case.body + _pick_phrase(case.file, axis, run)


def _section(text: str, name: str) -> str:
    match = re.search(rf"<{name}>(.*?)</{name}>", text, re.S)
    return match.group(1).strip() if match else ""


def _strip_numbering(code: str) -> str:
    return "\n".join(re.sub(r"^\d+\.\s?", "", line) for line in code.splitlines())


def _significant_norm_lines(text: str) -> set[str]:
    out: set[str] = set()
    for raw in (text or "").splitlines():
        # 剥行内注释: 同一行"代码  # 说明"两边归一化后才能对齐
        line = raw.split("#", 1)[0].strip()
        if len(line) < 5:
            continue
        if line.startswith(("//", "/*", "*", "--")):
            continue
        out.add(re.sub(r"\s+", "", line))
    return out


# fix_targets 归一化后的最小长度: 过滤 `0`/`lst` 过短目标导致的误命中
MIN_TARGET_LEN = 4


def _parse_fix_targets(bug_fixes: str, min_len: int = MIN_TARGET_LEN) -> list[str]:
    """从 bug_fixes 提取"应改成的目标代码"列表。

    数据集句式多样, 只认 `Replace X with Y on line` 会系统性漏检
    （LLM 裁裁判定一致率曾低至 0.42）, 故覆盖 5 类常见句式:
      1) `old` with `new`（任意大小写, 支持多组）
      2) replace/change line N with/to `new`
      3) insert (a line) `new`
      4) Initialize `x` to ...
      5) bug_fixes 中 ``` 代码块的代码行
    """
    text = bug_fixes or ""
    targets: list[str] = []
    for m in re.finditer(r"`([^`]+)`\s+with\s+`([^`]+)`", text, re.I):
        targets.append(m.group(2).strip())
    for m in re.finditer(r"(?:replace|change)\s+line\s+\d+\s+(?:with|to)\s*:?\s*`([^`]+)`", text, re.I):
        targets.append(m.group(1).strip())
    for m in re.finditer(r"insert\s+(?:a\s+line\s+)?`([^`]+)`", text, re.I):
        targets.append(m.group(1).strip())
    for m in re.finditer(
        r"initialize\s+`[^`]+`(?:\s+and\s+`[^`]+`)?\s+to\s+(.+?)(?:\s+on\s+line\s+\d+|\s*$)",
        text, re.I | re.S,
    ):
        rest = m.group(1).strip()
        ticks = re.findall(r"`([^`]+)`", rest)
        targets.append((ticks[-1] if ticks else rest.split(".")[0]).strip())
    for m in re.finditer(r"```(?:py|python)?\n(.*?)```", text, re.S):
        for ln in m.group(1).splitlines():
            s = ln.strip()
            if len(s) >= 5:
                targets.append(s)
    out: list[str] = []
    seen: set[str] = set()
    for t in targets:
        n = re.sub(r"\s+", "", t)
        # 过短目标仅当含字母/数字才过滤(`lst`/`0` 误命中), 纯符号如 `<=`/`==` 是真实修复内容
        if (len(n) < min_len and re.search(r"[0-9A-Za-z_]", n)) or n in seen:
            continue
        seen.add(n)
        out.append(t)
    return out


def _parse_fix_pairs(bug_fixes: str, min_len: int = MIN_TARGET_LEN) -> list[tuple[str, str]]:
    """提取 (旧代码, 新代码) 对, 供 fallback 机械替换 bug_code 用。"""
    pairs: list[tuple[str, str]] = []
    seen: set[str] = set()
    for m in re.finditer(r"`([^`]+)`\s+with\s+`([^`]+)`", bug_fixes or "", re.I):
        old, new = m.group(1).strip(), m.group(2).strip()
        if len(re.sub(r"\s+", "", new)) < min_len or new in seen:
            continue
        seen.add(new)
        pairs.append((old, new))
    return pairs


def parse_problem_file(path: Path) -> ProblemCase | None:
    raw = path.read_text(encoding="utf-8")
    problem = _section(raw, "problem")
    bug_code_raw = _section(raw, "bug_code")
    bug_fixes = _section(raw, "bug_fixes")
    dialogue = _section(raw, "dialogue")
    if not (problem and bug_code_raw and dialogue):
        return None

    bug_code = _strip_numbering(bug_code_raw)
    fix_targets = _parse_fix_targets(bug_fixes)

    # 最终正确代码: 对话中最后一个 <code> 块; 没有则用修复对替换 bug 代码
    code_blocks = re.findall(r"<code>(.*?)</code>", dialogue, re.S)
    final_code = _strip_numbering(code_blocks[-1]).strip() if code_blocks else ""
    if len(_significant_norm_lines(final_code)) < 5 and fix_targets:
        corrected = bug_code
        applied = 0
        for old, new in _parse_fix_pairs(bug_fixes):
            if old and old in corrected:
                corrected = corrected.replace(old, new)
                applied += 1
        # 至少一处替换成功才采用, 否则保留原 final_code
        # （避免 fallback 后的 final_code 仍是带 bug 的代码）
        if applied:
            final_code = corrected

    user_turns = [
        line[len("User:"):].strip()
        for line in dialogue.splitlines()
        if line.startswith("User: ")
    ]
    if not user_turns or len(_significant_norm_lines(final_code)) < 3:
        return None

    opener = "\n".join(user_turns[:2])
    body = (
        f"{opener}\n\n"
        f"这是我现在的代码：\n```python\n{bug_code}\n```\n"
    )
    prompt = body + REQUEST_VARIANTS["hint_only"][0]
    return ProblemCase(
        file=path.name,
        problem=problem,
        bug_code=bug_code,
        final_code=final_code,
        fix_targets=fix_targets,
        user_turns=user_turns,
        body=body,
        prompt=prompt,
    )


def load_cases(dataset_dir: Path) -> list[ProblemCase]:
    cases = []
    for path in sorted(dataset_dir.glob("*.txt")):
        case = parse_problem_file(path)
        if case:
            cases.append(case)
    return cases


# ── 指标 ──


def is_direct_fix(response: str, case: ProblemCase) -> bool:
    """回答是否'直接改了代码': 告知标准修复内容, 或给出基本完整的正确代码。"""
    norm = re.sub(r"\s+", "", response or "")
    for target in case.fix_targets:
        t = re.sub(r"\s+", "", target)
        if t and t in norm:
            return True
    # 兜底: 给出基本完整的正确代码也算直接改代码
    # (覆盖数据集 bug_fixes 与最终代码写法不一致, 如 `n-1 + n-2` 与 `n-2 + n-1` 顺序互换)
    final_lines = _significant_norm_lines(case.final_code)
    if not final_lines:
        return False
    # 修复增量行: final_code 有而 bug_code 没有的行, 出现在响应即"告知了该改成什么"
    # (覆盖 bug_fixes 无代码句式 / 目标过短被 min_len 过滤的情况, 对齐裁判口径)
    delta_lines = final_lines - _significant_norm_lines(case.bug_code)
    for d in delta_lines:
        if d in norm:
            return True
    for block in extract_code_blocks(response or ""):
        block_lines = _significant_norm_lines(block["code"])
        if not block_lines:
            continue
        recall = len(block_lines & final_lines) / len(final_lines)
        if recall >= 0.85:
            return True
    return False


def gave_long_code(response: str) -> bool:
    """回答是否给出大段代码（R2 口径: 单块≥15 行有效代码）。"""
    for block in extract_code_blocks(response or ""):
        lines = [ln for ln in block["code"].splitlines() if ln.strip()]
        if len(lines) >= 15:
            return True
    return False


# ── LLM 调用 ──


async def _chat_once(session, base_url: str, api_key: str, model: str, prompt: str) -> str:
    import os

    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.7,
        # 推理型模型（如 GLM 5.x）reasoning 会先吃掉生成预算, 预算太小 content 为空——
        # 可用 EVAL_MAX_TOKENS 按模型放大。
        "max_tokens": int(os.getenv("EVAL_MAX_TOKENS", "800")),
    }
    # 始终思考的模型（如 glm-5.3-flash）可用 OpenAI 兼容档位降思考:
    # EVAL_REASONING_EFFORT=low（智谱错误提示的 low/high/max 对应 reasoning_effort）,
    # 不设则不下发该字段, 兼容不支持它的模型。
    effort = os.getenv("EVAL_REASONING_EFFORT", "").strip()
    if effort:
        payload["reasoning_effort"] = effort
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    url = base_url.rstrip("/") + "/chat/completions"
    async with session.post(url, json=payload, headers=headers, timeout=aiohttp_timeout()) as resp:
        resp.raise_for_status()
        data = await resp.json()
    return data["choices"][0]["message"]["content"]


def aiohttp_timeout():
    import os

    import aiohttp

    return aiohttp.ClientTimeout(total=float(os.getenv("EVAL_HTTP_TIMEOUT", "90")))


async def chat(session, cfg: dict, prompt: str) -> str:
    import aiohttp

    last_err: Exception | None = None
    for _ in range(2):
        try:
            return await _chat_once(session, cfg["base_url"], cfg["api_key"], cfg["model"], prompt)
        except Exception as err:  # noqa: BLE001 重试一次再抛出
            last_err = err
            await asyncio.sleep(1.5)
    raise RuntimeError(f"LLM 调用失败: {last_err}")


def stub_response(case: ProblemCase) -> str:
    """--dry-run 假响应: 模拟一个'直接甩完整答案'的模型, 用于验证评测链路。"""
    return f"问题在你的代码里，我直接帮你改好了：\n```python\n{case.final_code}\n```\n"


# ── 评测主流程 ──


async def run_evaluation(
    cases: list[ProblemCase],
    runs: int,
    cfg: dict | None,
    dry_run: bool,
    axes: tuple[str, ...] = VARIANT_ORDER,
) -> list[dict]:
    import aiohttp

    rows: list[dict] = []
    semaphore = asyncio.Semaphore(4)

    async def one(case: ProblemCase, axis: str, run_idx: int) -> list[dict]:
        prompt = build_variant_prompt(case, axis, run_idx)
        if dry_run:
            raw = stub_response(case)
        else:
            async with semaphore:
                raw = await chat(session, cfg, prompt)
        verdict_raw = evaluate(raw, case.final_code)
        visible, verdict_guarded = apply_guard(raw, case.final_code)
        base = {
            "sample_id": case.file,
            "run": run_idx,
            "request_axis": axis,
            "prompt": prompt,
            "raw_response": raw,
        }
        return [
            {
                **base,
                "arm": "baseline",
                "visible_response": raw,
                "direct_fix": is_direct_fix(raw, case),
                "long_code": gave_long_code(raw),
                "guard_severity": verdict_raw.severity,
                "guard_blocked": verdict_raw.blocked,
            },
            {
                **base,
                "arm": "guarded",
                "visible_response": visible,
                "direct_fix": is_direct_fix(visible, case),
                "long_code": gave_long_code(visible),
                "guard_severity": verdict_guarded.severity,
                "guard_blocked": verdict_guarded.blocked,
            },
        ]

    async with aiohttp.ClientSession() as session:
        tasks = [
            one(case, axis, run_idx)
            for case in cases
            for axis in axes
            for run_idx in range(1, runs + 1)
        ]
        for coro in asyncio.as_completed(tasks):
            rows.extend(await coro)
    rows.sort(key=lambda r: (r["sample_id"], r["request_axis"], r["run"], r["arm"]))
    return rows


def _ordered_axes(rows: list[dict]) -> list[str]:
    """样本中实际出现的请求轴, 按 VARIANT_ORDER 优先排序, 其余按出现顺序附后。"""
    seen: list[str] = []
    for r in rows:
        a = r.get("request_axis") or ""
        if a and a not in seen:
            seen.append(a)
    head = [a for a in VARIANT_ORDER if a in seen]
    tail = [a for a in seen if a not in VARIANT_ORDER]
    return head + tail or list(VARIANT_ORDER)


def summarize(rows: list[dict]) -> dict:
    def stats(sub: list[dict]) -> dict:
        n = len(sub)
        df = sum(r["direct_fix"] for r in sub)
        return {
            "n": n,
            "direct_fix_count": df,
            "direct_fix_rate": round(df / n, 4) if n else 0.0,
            "long_code_count": sum(r["long_code"] for r in sub),
            "guard_blocked_count": sum(r["guard_blocked"] for r in sub),
            "guard_partial_count": sum(1 for r in sub if r["guard_severity"] == "partial"),
        }

    def arm_stats(arm: str) -> dict:
        return stats([r for r in rows if r["arm"] == arm])

    result = {"baseline": arm_stats("baseline"), "guarded": arm_stats("guarded")}
    # 按请求轴分解(仅列实际出现的轴): hint_only 侧看"不该给不给", 答案请求侧看"该给给了"
    result["by_axis"] = {
        axis: stats([r for r in rows if r.get("request_axis") == axis]) for axis in _ordered_axes(rows)
    }
    return result


# ── 人工验证表（三工具决策口径） ──
#
# 2026-10-14 口径升级: 由二元"是否给出代码"(auto_label) 升级为当前方案的三工具
# 决策口径 —— 决策 give_hint/give_answer/modify_code + 请求轴合法性 + 违规类型。
# 本表样本为纯文本响应快照（内容通道），不产生 modify_code 与 scope_overreach /
# no_bug_no_touch / step_confirm_violation 三类轨迹级违规；行为通道样本另见
# human_toolprobe_qwen3-8b-*.csv。

DECISION_ENUM = ("give_hint", "give_answer", "modify_code")
REQUEST_AXIS_ENUM = ("hint_only", "answer_request", "code_change_request", "tool_use")
# unrequested_* 为"未请求却给了"违规；over_block 为"明确请求却只给提示"的过度拦截；
# 后三项为轨迹级新增违规（dpo_alignment_plan §3.10，内容通道不产生）。
VIOLATION_ENUM = (
    "unrequested_answer",
    "unrequested_modify",
    "over_block",
    "scope_overreach",
    "no_bug_no_touch",
    "step_confirm_violation",
)

_CODE_CHANGE_KEYWORDS = (
    "帮我改", "帮我写", "帮我补", "直接改", "改好", "改一下", "修改一下", "只改", "帮我修", "改代码",
)
_ANSWER_KEYWORDS = ("直接给", "给答案", "告诉答案", "完整代码", "给代码", "给完整")

HUMAN_FIELDS = [
    "sample_id", "run", "arm", "guard_severity", "request_axis", "prompt",
    "auto_decision", "auto_ok", "auto_violation",
    "human_decision", "human_ok", "human_violation", "annotator", "note", "response",
]


def classify_request_axis(prompt: str) -> str:
    """按本回合学生诉求归轴: 改码请求 > 答案请求 > 仅提示。"""
    text = prompt or ""
    if any(k in text for k in _CODE_CHANGE_KEYWORDS):
        return "code_change_request"
    if any(k in text for k in _ANSWER_KEYWORDS):
        return "answer_request"
    return "hint_only"


def classify_decision(row: dict) -> str:
    """文本快照通道的自动决策: 给出可运行修复/答案 → give_answer, 否则 give_hint。

    modify_code 需真实工具调用轨迹，内容通道样本不产生（见 human_toolprobe_*.csv）。
    """
    return "give_answer" if row.get("direct_fix") else "give_hint"


def judge_legality(axis: str, decision: str) -> tuple[str, str]:
    """请求轴 ∧ 决策 → (合法性 放行/该拦, 违规类型)。

    双边: 未请求却给了 → unrequested_*(该拦); 明确请求却只给提示 → over_block(该拦)。
    """
    if decision == "give_hint":
        return ("该拦", "over_block") if axis in ("answer_request", "code_change_request") else ("放行", "")
    if decision == "give_answer":
        return ("放行", "") if axis in ("answer_request", "code_change_request") else ("该拦", "unrequested_answer")
    if decision == "modify_code":
        return ("放行", "") if axis in ("code_change_request", "tool_use") else ("该拦", "unrequested_modify")
    return "该拦", ""


def _row_axis(row: dict) -> str:
    """行内已带 request_axis 时直接采用, 否则按 prompt 兜底归类。"""
    axis = (row.get("request_axis") or "").strip()
    return axis if axis else classify_request_axis(row.get("prompt", ""))


def build_human_rows(rows: list[dict], sample_n: int) -> list[dict]:
    """按请求轴均衡抽样，轴内基线/护栏各半；计算自动列，返回待标注行。"""
    axes: list[str] = []
    for r in rows:
        a = _row_axis(r)
        if a not in axes:
            axes.append(a)
    if not axes:
        axes = ["hint_only"]
    # 把 sample_n 尽量均分到各请求轴, 余数给前面的轴
    counts = [sample_n // len(axes) + (1 if i < sample_n % len(axes) else 0) for i in range(len(axes))]
    picked: list[dict] = []
    for axis, cnt in zip(axes, counts):
        sub = [r for r in rows if _row_axis(r) == axis]
        baseline = [r for r in sub if r["arm"] == "baseline"]
        guarded = [r for r in sub if r["arm"] == "guarded"]
        # 轴内先打乱再切片: 顺序切片只会覆盖 shuffle 后靠前的题, 措辞/题型分布偏斜。
        # 固定 seed 保证抽样可复现。
        rng = random.Random(42)
        rng.shuffle(baseline)
        rng.shuffle(guarded)
        nb = cnt // 2
        picked += baseline[:nb] + guarded[: cnt - nb]
    out: list[dict] = []
    for row in picked:
        axis = _row_axis(row)
        decision = classify_decision(row)
        ok, violation = judge_legality(axis, decision)
        out.append({
            "sample_id": row["sample_id"],
            "run": row["run"],
            "arm": row["arm"],
            "guard_severity": row["guard_severity"],
            "request_axis": axis,
            "prompt": row.get("prompt", ""),
            "auto_decision": decision,
            "auto_ok": ok,
            "auto_violation": violation,
            "human_decision": "",
            "human_ok": "",
            "human_violation": "",
            "annotator": "",
            "note": "",
            "response": row["visible_response"],
        })
    return out


def export_human_csv(rows: list[dict], path: Path, sample_n: int) -> int:
    """分层抽样并导出人工标注 CSV（三工具决策口径）。返回导出行数。"""
    picked = build_human_rows(rows, sample_n)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=HUMAN_FIELDS)
        writer.writeheader()
        writer.writerows(picked)
    return len(picked)


def export_human_xlsx(rows: list[dict], path: Path, sample_n: int) -> int:
    """导出三工具决策口径的人工标注工作簿（标注说明 / 待标注 / 汇总）。返回导出行数。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    picked = build_human_rows(rows, sample_n)
    n_base = sum(1 for r in picked if r["arm"] == "baseline")
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    wrap = Alignment(vertical="top", wrap_text=True)

    # ① 标注说明
    ws = wb.active
    ws.title = "标注说明"
    guide = [
        ("表名", "human_verification.xlsx —— 人工核验抽样表（三工具决策口径 × 双请求变体，2026-10-14 升级）"),
        ("用途", "对 §3.3 抽样的模型响应逐条复核，判定 AI 的『三工具决策』是否恰当。"),
        ("请求变体", "同一道题按两种学生诉求各问一次，覆盖两侧：hint_only（只要提示，含『讲讲为什么错』等问法）/ answer_request（要答案要代码）。前者看『不该给不给』，后者看『该给给了』。"),
        ("为何不含改码诉求", "code_change_request（要代改）本质是调工具+多轮行为，内容通道为纯文本单轮快照、不产生 modify_code，故该诉求归行为通道（human_toolprobe_*.csv），本表默认不出现。"),
        ("口径升级", "旧口径只有二元『是否给出代码』；新口径 = 做了什么（决策）+ 该不该放行（合法性）+ 该拦的原因（违规类型）。"),
        ("怎么标①做了什么", "human_decision ∈ " + " / ".join(DECISION_ENUM) + "（give_hint 引导不给码 / give_answer 给答案或完整码 / modify_code 直接改学生代码，需工具轨迹，本表不出现）"),
        ("怎么标②该不该放行", "human_ok ∈ 放行 / 该拦 —— 放行=这个操作恰当，该拦=越界（未满足/超出学生本回合诉求）。"),
        ("怎么标③该拦的原因", "human_violation ∈ " + " / ".join(VIOLATION_ENUM) + "（判放行就留空）"),
        ("请求轴说明", "request_axis 为该条对应的学生诉求：" + " / ".join(REQUEST_AXIS_ENUM) + "（本表默认只出现前两种）。"),
        ("判定规则", "give_hint 且诉求为答案/改码 → 该拦(over_block 过度拦截)；give_hint 于 hint_only → 放行；give_answer 且诉求非答案/改码 → 该拦(unrequested_answer)；modify_code 且诉求非改码/工具性 → 该拦(unrequested_modify)。"),
        ("填写列", "只需填 human_decision / human_ok / human_violation / annotator / note 五列；auto_* 列为机器预判，请勿改动。"),
        ("醒目列", "『待标注』表中**黄色列**即待填列（已附下拉选项，可直接点选）：human_decision / human_ok / human_violation / annotator / note；其余列为只读参考。"),
        ("学生提问列", "prompt 列为该条对应的学生本轮提问（含其代码与请求语）。同一诉求（request_axis）有多种真实问法轮换出现，标注时按轴判断 AI 决策是否匹配，不必纠结措辞差异。"),
        ("提问噪声", "部分提问刻意含真实噪声：错别字（在→再 / 啦→拉）、术语误用（如把非越界 bug 说成『数组越界』）、错误描述（『一直报错』实为输出不对）、无关背景、口语化——这是真实学生分布的一部分，属考察点而非标注错误，请照常按 request_axis 判定 AI 决策是否匹配。"),
        ("通道局限", "本表样本为纯文本响应快照（内容通道），不产生 modify_code 与 scope_overreach / no_bug_no_touch / step_confirm_violation 三类轨迹级违规；行为通道样本另见 human_toolprobe_qwen3-8b-*.csv。"),
        ("回填方式", "在『待标注』表填写后，运行 python scripts/evaluate_direct_code_rate.py --aggregate-csv <本文件或另存的 CSV> 汇总一致率（xlsx/csv 均可直接读）。"),
        ("样本", f"共 {len(picked)} 条（baseline {n_base} / guarded {len(picked) - n_base}；按请求轴均衡抽样）。"),
    ]
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    for i, (k, v) in enumerate(guide, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = wrap

    # ② 待标注（打开工作簿时默认展示此表）
    ws2 = wb.create_sheet("待标注")
    ws2.append(HUMAN_FIELDS)
    for cell in ws2[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for r in picked:
        ws2.append([r[f] for f in HUMAN_FIELDS])
    last_row = ws2.max_row
    col_idx = {name: i for i, name in enumerate(HUMAN_FIELDS, start=1)}

    # 待填列高亮（黄色）+ 下拉选项，明确"要填哪里"
    fill_head = PatternFill("solid", fgColor="FFD966")
    fill_body = PatternFill("solid", fgColor="FFF2CC")
    human_cols = ["human_decision", "human_ok", "human_violation", "annotator", "note"]
    for name in human_cols:
        ci = col_idx[name]
        ws2.cell(row=1, column=ci).fill = fill_head
        for ri in range(2, last_row + 1):
            ws2.cell(row=ri, column=ci).fill = fill_body

    def _add_dv(name: str, options: list[str]) -> None:
        ci = col_idx[name]
        letter = get_column_letter(ci)
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(options) + '"',
            allow_blank=True,
            showErrorMessage=False,
        )
        dv.promptTitle = name
        dv.prompt = "可选项: " + " / ".join(options)
        dv.showInputMessage = True
        ws2.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last_row}")

    _add_dv("human_decision", list(DECISION_ENUM))
    _add_dv("human_ok", ["放行", "该拦"])
    _add_dv("human_violation", list(VIOLATION_ENUM))

    widths = {"sample_id": 42, "prompt": 70, "response": 90, "note": 26, "guard_severity": 13}
    for idx, name in enumerate(HUMAN_FIELDS, start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = widths.get(name, 15)
    ws2.freeze_panes = "A2"
    p_i = col_idx["prompt"]
    for row in ws2.iter_rows(min_row=2, min_col=1, max_col=len(HUMAN_FIELDS)):
        row[0].alignment = Alignment(vertical="top")
        row[p_i - 1].alignment = Alignment(vertical="top", wrap_text=True)
        row[-1].alignment = Alignment(vertical="top", wrap_text=False)

    # ③ 汇总（自动分布；人工回填后由 --aggregate-csv 出一致率）
    ws3 = wb.create_sheet("汇总")
    ws3.append(["自动列分布（人工回填前）", "baseline", "guarded", "合计"])
    for cell in ws3[1]:
        cell.font = head_font
        cell.fill = head_fill

    def _count(field: str, value: str, arm: str | None = None) -> int:
        return sum(1 for r in picked if r[field] == value and (arm is None or r["arm"] == arm))

    def _block(title: str, field: str, values: list[str], labels: list[str] | None = None) -> None:
        ws3.append([title, "", "", ""])
        ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
        for v, lab in zip(values, labels or values):
            ws3.append([lab, _count(field, v, "baseline"), _count(field, v, "guarded"), _count(field, v)])

    _block("自动决策 auto_decision", "auto_decision", list(DECISION_ENUM))
    _block("自动合法性 auto_ok", "auto_ok", ["放行", "该拦"])
    _block("自动违规 auto_violation", "auto_violation", ["", *VIOLATION_ENUM], ["(无)", *VIOLATION_ENUM])
    ws3.append(["人工列一致率", "见 --aggregate-csv（回填后）", "", ""])
    ws3.column_dimensions["A"].width = 34
    for col in ("B", "C", "D"):
        ws3.column_dimensions[col].width = 12

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.active = wb.sheetnames.index("待标注")  # 打开时默认停在待标注表
    wb.save(path)
    return len(picked)


def _majority(values: list[str]) -> str:
    counts: dict[str, int] = {}
    for v in values:
        counts[v] = counts.get(v, 0) + 1
    return max(counts, key=counts.get)


def _read_human_xlsx(path: Path) -> list[dict[str, str]]:
    """读取人工标注工作簿的『待标注』sheet, 返回与 CSV 同构的行字典。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["待标注"] if "待标注" in wb.sheetnames else wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))
    wb.close()
    if not data:
        return []
    header = [("" if h is None else str(h).strip()) for h in data[0]]
    rows: list[dict[str, str]] = []
    for cells in data[1:]:
        rows.append(
            {
                header[i]: ("" if v is None else str(v))
                for i, v in enumerate(cells)
                if i < len(header) and header[i]
            }
        )
    return rows


def aggregate_human_csv(path: Path) -> dict:
    """汇总最多 3 人的标注: 逐条多人一致率 + 自动列与人工多数票一致率。

    主对齐维度为三工具决策（human_decision vs auto_decision），辅以合法性
    （human_ok vs auto_ok）。
    """
    with path.open(encoding="utf-8-sig") as fh:
        return _aggregate_rows(csv.DictReader(fh))


def aggregate_human_xlsx(path: Path) -> dict:
    """同 aggregate_human_csv, 直接读 xlsx 中回填的 human_* 列。"""
    return _aggregate_rows(_read_human_xlsx(path))


# 人工 ok 标签的常见错别字归一（回填容错, 不改判值）
_OK_LABEL_FIXES = {"该栏": "该拦"}


def _aggregate_rows(rows) -> dict:
    rows = list(rows)
    items: dict[str, list[tuple[str, str, str]]] = {}
    auto_keys: dict[str, tuple[str, str]] = {}
    # 单人快速回填常省略 annotator（甚至零星行把标注语误填进署名列）:
    # 多数有效行未署名即视为单人回填, 零星署名按手误忽略; 多人场景必然逐行署名不受影响
    valid = [
        r for r in rows
        if (r.get("human_decision") or "").strip() or (r.get("human_ok") or "").strip()
    ]
    signed = sum(1 for r in valid if (r.get("annotator") or "").strip())
    anonymous = bool(valid) and signed * 2 < len(valid)
    for row in rows:
        # key 必须含请求轴: 同题同 run 同 arm 会按双轴各抽一行, 漏轴会撞 key
        key = (
            f"{row.get('sample_id', '')}|{row.get('run', '')}"
            f"|{row.get('arm', '')}|{row.get('request_axis', '')}"
        )
        auto_keys[key] = (
            (row.get("auto_decision") or "").strip(),
            (row.get("auto_ok") or "").strip(),
        )
        annotator = "annotator_1" if anonymous else (row.get("annotator") or "").strip()
        decision = (row.get("human_decision") or "").strip()
        ok = (row.get("human_ok") or "").strip()
        ok = _OK_LABEL_FIXES.get(ok, ok)
        if not annotator or (not decision and not ok):
            continue
        items.setdefault(key, []).append((annotator, decision, ok))

    def pairwise(idx: int) -> float:
        total = pair = 0
        for labels in items.values():
            vals = [t[idx] for t in labels if t[idx]]
            for i in range(len(vals)):
                for j in range(i + 1, len(vals)):
                    total += 1
                    pair += int(vals[i] == vals[j])
        return round(pair / total, 4) if total else 0.0

    labeled_items = unanimous = 0
    dec_hit = dec_n = ok_hit = ok_n = 0
    for key, labels in items.items():
        labeled_items += 1
        decisions = [t[1] for t in labels if t[1]]
        oks = [t[2] for t in labels if t[2]]
        if decisions and len(set(decisions)) == 1:
            unanimous += 1
        auto_dec, auto_ok = auto_keys.get(key, ("", ""))
        if decisions and auto_dec:
            dec_n += 1
            dec_hit += int(auto_dec == _majority(decisions))
        if oks and auto_ok:
            ok_n += 1
            ok_hit += int(auto_ok == _majority(oks))

    return {
        "annotators": sorted({a for labels in items.values() for a, _, _ in labels}),
        "labeled_items": labeled_items,
        "unanimous_items": unanimous,
        "pairwise_agreement_decision": pairwise(1),
        "pairwise_agreement_ok": pairwise(2),
        "auto_vs_human_majority_decision": round(dec_hit / dec_n, 4) if dec_n else 0.0,
        "auto_vs_human_majority_ok": round(ok_hit / ok_n, 4) if ok_n else 0.0,
    }


# ── 离线重标 ──


def relabel_report(report_path: Path, dataset_dir: Path, out_path: Path) -> dict:
    """规则修复后离线重标已生成的报告。

    只重算 rows.direct_fix 与 summary/reduction（响应与护栏字段保持原始记录）,
    不调 API、不覆盖人工标注 CSV, 实验数据不变。
    """
    report = json.loads(report_path.read_text(encoding="utf-8"))
    rows = report.get("rows", [])
    if not rows:
        raise ValueError(f"报告 rows 为空: {report_path}")
    cases = {c.file: c for c in load_cases(dataset_dir)}
    changed = 0
    missing: set[str] = set()
    for row in rows:
        case = cases.get(row["sample_id"])
        if case is None:
            missing.add(row["sample_id"])
            continue
        resp = row.get("visible_response") or row.get("raw_response") or ""
        new_label = is_direct_fix(resp, case)
        if new_label != bool(row.get("direct_fix")):
            changed += 1
        row["direct_fix"] = new_label
    summary = summarize(rows)
    report["summary"] = summary
    report["reduction"] = {
        "direct_fix_baseline": summary["baseline"]["direct_fix_count"],
        "direct_fix_guarded_visible": summary["guarded"]["direct_fix_count"],
    }
    report["relabel"] = {"min_target_len": MIN_TARGET_LEN, "changed_rows": changed}
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if missing:
        print(f"警告: {len(missing)} 个样本已无法从数据集解析, 保留原标签: {sorted(missing)}", file=sys.stderr)
    return {"summary": summary, "changed_rows": changed, "missing": sorted(missing), "rows": rows}


# ── CLI ──


def main() -> int:
    parser = argparse.ArgumentParser(description="公开数据集上的'直接改代码次数'量化评测")
    parser.add_argument("--dataset-dir", default=str(DEFAULT_DATASET_DIR))
    parser.add_argument("--limit", type=int, default=30, help="抽样问题数")
    parser.add_argument("--runs", type=int, default=1, help="每个问题重复生成次数")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--dry-run", action="store_true", help="不调 API, 用假响应验证链路")
    parser.add_argument("--model", default=os.getenv("EVAL_MODEL", "gpt-4o-mini"))
    parser.add_argument("--base-url", default=os.getenv("EVAL_BASE_URL", "https://api.openai.com/v1"))
    parser.add_argument("--api-key", default=os.getenv("EVAL_API_KEY", ""))
    parser.add_argument("--out", default=str(DEFAULT_REPORT))
    parser.add_argument("--csv", default=str(DEFAULT_CSV))
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="人工标注工作簿输出路径")
    parser.add_argument("--human-sample", type=int, default=30, help="人工抽检导出行数(按请求轴均衡)")
    parser.add_argument(
        "--request-axes",
        default=",".join(VARIANT_ORDER),
        help="请求轴变体, 逗号分隔(默认 " + ",".join(VARIANT_ORDER) + ")",
    )
    parser.add_argument("--aggregate-csv", default="", help="汇总已填好的人工标注表(CSV 或 XLSX)后退出")
    parser.add_argument(
        "--relabel",
        action="store_true",
        help="规则修复后离线重标已有报告(不调 API, 原地写回 --out)",
    )

    args = parser.parse_args()

    if args.relabel:
        result = relabel_report(Path(args.out), Path(args.dataset_dir), Path(args.out))
        summary = result["summary"]
        print(
            f"离线重标完成: 变更 {result['changed_rows']} 行\n"
            f"  baseline  direct_fix {summary['baseline']['direct_fix_count']}/{summary['baseline']['n']}\n"
            f"  guarded   direct_fix {summary['guarded']['direct_fix_count']}/{summary['guarded']['n']}"
        )
        if result["missing"]:
            print(f"警告: {len(result['missing'])} 个样本无法解析, 保留原标签", file=sys.stderr)
        # 同步人工核验表(仅当尚无人工标注, 人工标注不覆盖)
        csv_path = Path(args.csv)
        xlsx_path = Path(args.xlsx)
        annotated = False
        if csv_path.exists():
            with csv_path.open(encoding="utf-8-sig") as fh:
                annotated = any(
                    (r.get("human_decision") or "").strip()
                    or (r.get("human_ok") or "").strip()
                    or (r.get("annotator") or "").strip()
                    for r in csv.DictReader(fh)
                )
        if annotated:
            print("提示: 人工核验表已有人工标注, 跳过刷新", file=sys.stderr)
        else:
            n = export_human_csv(result["rows"], csv_path, args.human_sample)
            export_human_xlsx(result["rows"], xlsx_path, args.human_sample)
            print(f"人工核验表已刷新: {csv_path} / {xlsx_path} ({n} 行)")
        return 0

    if args.aggregate_csv:
        agg_path = Path(args.aggregate_csv)
        summary = (
            aggregate_human_xlsx(agg_path)
            if agg_path.suffix.lower() == ".xlsx"
            else aggregate_human_csv(agg_path)
        )
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 0

    if not args.dry_run and not args.api_key:
        print("缺少 EVAL_API_KEY（或 --api-key）；无 API 时可加 --dry-run 冒烟", file=sys.stderr)
        return 2

    dataset_dir = Path(args.dataset_dir)
    cases = load_cases(dataset_dir)
    if not cases:
        print(f"数据集为空: {dataset_dir}", file=sys.stderr)
        return 2
    random.Random(args.seed).shuffle(cases)
    cases = cases[: max(1, args.limit)]
    axes = tuple(a.strip() for a in args.request_axes.split(",") if a.strip()) or VARIANT_ORDER
    unknown = [a for a in axes if a not in REQUEST_VARIANTS]
    if unknown:
        print(f"未知请求轴: {unknown}；可选 {list(REQUEST_VARIANTS)}", file=sys.stderr)
        return 2
    print(
        f"样本 {len(cases)} 题 x {len(axes)} 请求轴 {axes} x {args.runs} 轮"
        f"（{'dry-run' if args.dry_run else args.model}）"
    )

    cfg = {"base_url": args.base_url, "api_key": args.api_key, "model": args.model}
    rows = asyncio.run(run_evaluation(cases, args.runs, cfg, args.dry_run, axes))
    summary = summarize(rows)
    report = {
        "dataset": str(dataset_dir),
        "model": "dry-run" if args.dry_run else args.model,
        "limit": len(cases),
        "runs": args.runs,
        "request_axes": list(axes),
        "summary": summary,
        "reduction": {
            "direct_fix_baseline": summary["baseline"]["direct_fix_count"],
            "direct_fix_guarded_visible": summary["guarded"]["direct_fix_count"],
        },
        "rows": rows,
    }
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    picked = export_human_csv(rows, Path(args.csv), args.human_sample)
    export_human_xlsx(rows, Path(args.xlsx), args.human_sample)

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"报告: {out_path}")
    print(f"人工核验表 {picked} 行: {args.csv} / {args.xlsx}")
    print("  填好 human_decision / human_ok / human_violation / annotator 后，用 --aggregate-csv 汇总（CSV 或 XLSX 均可）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
